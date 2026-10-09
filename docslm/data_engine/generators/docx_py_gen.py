"""DOCX edit/format/read routes via python-docx (attachment-style) — English-only."""
from string import Template
import json as _json
from .. import content as C

_SETUP = Template("""# setup: attachment base document
from docx import Document
from docx.shared import Pt

doc = Document()
doc.add_heading("$base_title", level=1)
for t in $paras:
    p = doc.add_paragraph(t)
    if $bad_spacing:
        p.paragraph_format.line_spacing = Pt(12)   # too tight on purpose
doc.add_heading("$h2", level=2)
for t in $paras2:
    p = doc.add_paragraph(t)
doc.add_table(rows=3, cols=3, style="Table Grid")
doc.save("$base")
print("ok")
""")

_EDIT = Template("""# DocSLM edit route — python-docx targeted edit
import json
from docx import Document

BASE, OUT = "$base", "$artifact"
doc = Document(BASE)

# 1. apply the requested content edit: retitle + append new section
doc.paragraphs[0].text = "$new_title"
doc.add_heading("$h2_new", level=2)
doc.add_paragraph("$new_para")

# 2. normalize formatting so the result passes postcheck: 1.3x spacing everywhere
def _fix(container):
    for p in container.paragraphs:
        p.paragraph_format.line_spacing = 1.3

_fix(doc)
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            _fix(cell)
doc.save(OUT)
print(json.dumps({"exit": 0, "artifact": OUT, "edited": True}))
""")

_FORMAT = Template("""# DocSLM format route — normalize spacing + heading styles
import json
from docx import Document

BASE, OUT = "$base", "$artifact"
doc = Document(BASE)

for p in doc.paragraphs:
    p.paragraph_format.line_spacing = 1.3
    if p.style is not None and p.style.name.startswith("Heading"):
        for r in p.runs:
            r.font.bold = True

def _fix(cell_or_doc):
    for p in cell_or_doc.paragraphs:
        p.paragraph_format.line_spacing = 1.3

for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            _fix(cell)
doc.save(OUT)
print(json.dumps({"exit": 0, "artifact": OUT, "formatted": True}))
""")

_READ = Template("""# DocSLM read route — extract structure and content
import json
from docx import Document

doc = Document("$base")
title = doc.paragraphs[0].text if doc.paragraphs else ""
body = [p.text for p in doc.paragraphs if p.text.strip()][1:11]
tables = [[[cell.text for cell in row.cells] for row in t.rows] for t in doc.tables]
out = {"title": title, "paragraph_count": len(doc.paragraphs),
       "body_sample": body, "tables": tables}
with open("$artifact", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print(json.dumps({"exit": 0, "artifact": "$artifact", "title": title}))
""")


def gen_edit(spec, rng):
    base, artifact = f"{spec.task_id}_base.docx", f"{spec.task_id}_edited.docx"
    interim_title = f"{C.project(rng, spec.lang)} Interim Report"
    new_title = f"{C.project(rng, spec.lang)} Closure Report"
    h2_new = "Acceptance Outcome"
    new_para = "All deliverables are accepted; the project enters the warranty period."
    setup = _SETUP.substitute(
        base_title=interim_title,
        paras=_json.dumps(["Progress on track; risks contained.",
                           "Milestones M1 and M2 delivered on time."], ensure_ascii=False),
        h2="Milestone Review",
        paras2=_json.dumps(["Integration testing completed."], ensure_ascii=False),
        bad_spacing="True", base=base)
    code = _EDIT.substitute(base=base, artifact=artifact,
                            new_title=new_title, h2_new=h2_new, new_para=new_para)
    request = (f"The attachment {base} is an interim project doc. Retitle it to "
               f"'{new_title}', append a '{h2_new}' section with the closing paragraph, "
               f"normalize 1.3x line spacing, save as {artifact}.")
    return {"runner": "python", "code": code, "setup_runner": "python", "setup_code": setup,
            "artifact": artifact, "expect": {"kind": "docx"}, "request": request, "mutation": None}


def gen_format(spec, rng):
    base = artifact = f"{spec.task_id}_fmt.docx"  # in-place reformat
    title = "Department Quarterly Summary"
    setup = _SETUP.substitute(
        base_title=title,
        paras=_json.dumps(C.report_sections(rng, spec.lang, 2)[0][1], ensure_ascii=False),
        h2="Data Analysis",
        paras2=_json.dumps([C.report_sections(rng, spec.lang, 2)[1][1]], ensure_ascii=False),
        bad_spacing="True", base=base)
    code = _FORMAT.substitute(base=base, artifact=artifact)
    request = (f"The attachment {base} has cramped spacing and inconsistent formatting. "
               f"Normalize to 1.3x line spacing, keep heading levels, save in place.")
    return {"runner": "python", "code": code, "setup_runner": "python", "setup_code": setup,
            "artifact": artifact, "expect": {"kind": "docx"}, "request": request, "mutation": None}


def gen_read(spec, rng):
    base, artifact = f"{spec.task_id}_src.docx", f"{spec.task_id}_extract.json"
    title = C.report_sections(rng, spec.lang, 2)[0][0]
    setup = _SETUP.substitute(
        base_title=title,
        paras=_json.dumps(["Body paragraph one.", "Body paragraph two."], ensure_ascii=False),
        h2="Appendix Table",
        paras2=_json.dumps(["Table notes."], ensure_ascii=False),
        bad_spacing="False", base=base)
    code = _READ.substitute(base=base, artifact=artifact)
    request = (f"Read {base}: extract title, paragraph stats, body sample and tables "
               f"as JSON to {artifact}.")
    return {"runner": "python", "code": code, "setup_runner": "python", "setup_code": setup,
            "artifact": artifact, "expect": {"kind": "extract"}, "request": request, "mutation": None}
