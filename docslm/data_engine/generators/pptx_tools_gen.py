"""PPTX inspect/edit generator — python-pptx edit verified structurally."""
from string import Template
import json as _json
from .. import content as C

_DECK_SETUP = Template("""// fixture: deck for inspection ($label)
const pptxgen = require("pptxgenjs");
const pptx = new pptxgen();
pptx.defineLayout({ name: "WIDE", width: 13.33, height: 7.5 });
pptx.layout = "WIDE";
const s1 = pptx.addSlide();
s1.addText("$title", { x: 0.8, y: 2.4, w: 11.7, h: 1.3, fontSize: 40, bold: true });
const s2 = pptx.addSlide();
s2.addText("Agenda", { x: 0.5, y: 0.3, w: 12, h: 0.9, fontSize: 24, bold: true });
s2.addText("Status, numbers, next steps.", { x: 0.6, y: 1.3, w: 12, h: 4, fontSize: 16 });
const s3 = pptx.addSlide();
s3.addText("Numbers", { x: 0.5, y: 0.3, w: 12, h: 0.9, fontSize: 24, bold: true });
pptx.writeFile({ fileName: "$name" }).then(() => console.log(JSON.stringify({ exit: 0 })));
""")

_INSPECT_EDIT = Template('''# DocSLM pptx inspect_edit — python-pptx edit
import json
from pptx import Presentation
from pptx.util import Pt

SRC, OUT = "$name", "$artifact"
prs = Presentation(SRC)
n_before = len(prs.slides)

slide = prs.slides.add_slide(prs.slide_layouts[len(prs.slide_layouts) - 1])
box = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(600), Pt(80))
tf = box.text_frame
tf.text = "$new_slide_title"
tf.paragraphs[0].runs[0].font.size = Pt(28)
tf.paragraphs[0].runs[0].font.bold = True

prs.save(OUT)
check = Presentation(OUT)
assert len(check.slides) == n_before + 1, "slide not appended"
texts = [sh.text_frame.text for s in check.slides for sh in s.shapes
         if sh.has_text_frame and sh.text_frame.text]
assert "$new_slide_title" in texts, "new slide title missing"
print(json.dumps({"exit": 0, "artifact": OUT, "slides": len(check.slides)}))
''')


def gen_inspect_edit(spec, rng):
    name = f"{spec.task_id}_deck.pptx"
    artifact = f"{spec.task_id}_upd.pptx"
    title = C.report_sections(rng, spec.lang, 1)[0][0]
    new_title = "Next Steps"
    setup = _DECK_SETUP.substitute(label="3-slide deck", name=name, title=title)
    code = _INSPECT_EDIT.substitute(name=name, artifact=artifact, new_slide_title=new_title)
    request = (f"Append a 'Next Steps' slide to {name} using python-pptx and save the "
               f"updated deck as {artifact}; verify the slide count increased.")
    return {"runner": "python", "code": code, "setup_runner": "node", "setup_code": setup,
            "artifact": artifact,
            "expect": {"kind": "pptx", "min_slides": 4}, "request": request, "mutation": None}
