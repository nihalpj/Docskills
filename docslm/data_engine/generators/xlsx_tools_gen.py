"""XLSX tool-route generators — operate the xlsx skill's own QA CLI
(xlsx.py inspect/scan/audit/pivot). English-only. recalc is excluded
(requires LibreOffice, not present in this sandbox).
"""
from string import Template
import json as _json
from .. import content as C

_XLSX = '"/skills/xlsx/xlsx.py"'

_BOOK_SETUP = Template("""# setup: workbook fixture ($label)
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "$sheet"
ws.append(["$dim", "$metric"])
$append_rows
wb.save("$name")
print("ok")
""")

_INSPECT = Template('''# DocSLM xlsx inspect — workbook structure report via the skill CLI
import json, subprocess, sys

proc = subprocess.run([sys.executable, $xlsx, "inspect", "$name"],
                      capture_output=True, text=True, timeout=120)
data = json.loads(proc.stdout)
sheets = data.get("sheets", [])
assert sheets and sheets[0]["name"] == "$sheet", f"unexpected inspect output: {proc.stdout[:200]}"
print(json.dumps({"exit": 0, "artifact": "$name", "sheets": [s["name"] for s in sheets]}))
''')

_SCAN = Template('''# DocSLM xlsx scan — issue scan via the skill CLI
import json, subprocess, sys

proc = subprocess.run([sys.executable, $xlsx, "scan", "$name"],
                      capture_output=True, text=True, timeout=120)
data = json.loads(proc.stdout)
assert "total_findings" in data, f"unexpected scan output: {proc.stdout[:200]}"
assert data["total_findings"] == 0, f"findings: {data}"
print(json.dumps({"exit": 0, "artifact": "$name", "findings": 0}))
''')

_AUDIT = Template('''# DocSLM xlsx audit — formula audit via the skill CLI
import json, subprocess, sys

proc = subprocess.run([sys.executable, $xlsx, "audit", "$name"],
                      capture_output=True, text=True, timeout=120)
data = json.loads(proc.stdout)
assert "total_formulas" in data, f"unexpected audit output: {proc.stdout[:200]}"
assert data.get("error_count", 0) == 0, f"formula errors: {data}"
print(json.dumps({"exit": 0, "artifact": "$name", "formulas": data["total_formulas"]}))
''')

_PIVOT = Template('''# DocSLM xlsx pivot — PivotTable summary via the skill CLI
import json, subprocess, sys

proc = subprocess.run([sys.executable, $xlsx, "pivot",
                       "--source", "$range", "--values", "$metric", "--rows", "$dim",
                       "$name", "$out"],
                      capture_output=True, text=True, timeout=120)
assert proc.returncode == 0, f"pivot failed: {(proc.stdout + proc.stderr)[-300:]}"
from openpyxl import load_workbook
wb = load_workbook("$out")
assert len(wb.sheetnames) >= 2, f"pivot output missing sheet: {wb.sheetnames}"
print(json.dumps({"exit": 0, "artifact": "$out", "sheets": wb.sheetnames}))
''')


def _book_setup(spec, rng, label):
    dim, metric = "Region", "Revenue"
    rows = _json.dumps([[r, rng.randrange(100, 900)] for r in
                        ["North", "South", "East", "West"]])
    return _BOOK_SETUP.substitute(label=label, sheet="Data", dim=dim, metric=metric,
                                  append_rows=f"for r in {rows}:\n    ws.append(r)\n",
                                  name=f"{spec.task_id}_wb.xlsx")


def gen_inspect(spec, rng):
    name = f"{spec.task_id}_wb.xlsx"
    setup = _book_setup(spec, rng, "workbook for inspection")
    code = _INSPECT.substitute(xlsx=_XLSX, name=name, sheet="Data")
    request = (f"Inspect the structure of {name} with the xlsx skill's inspect command "
               f"and report the sheet list and data range.")
    return {"runner": "python", "code": code, "setup_runner": "python", "setup_code": setup,
            "artifact": name, "expect": {"kind": "tool", "json_status": False,
                                         "artifact": name}, "request": request, "mutation": None}


def gen_scan(spec, rng):
    name = f"{spec.task_id}_wb.xlsx"
    setup = _book_setup(spec, rng, "clean workbook for scanning")
    code = _SCAN.substitute(xlsx=_XLSX, name=name)
    request = (f"Scan {name} for spreadsheet issues with the xlsx skill's scan command "
               f"and confirm no findings are reported.")
    return {"runner": "python", "code": code, "setup_runner": "python", "setup_code": setup,
            "artifact": name, "expect": {"kind": "tool", "json_status": False,
                                         "artifact": name}, "request": request, "mutation": None}


def gen_audit(spec, rng):
    name = f"{spec.task_id}_wb.xlsx"
    setup = _book_setup(spec, rng, "workbook for formula audit")
    code = _AUDIT.substitute(xlsx=_XLSX, name=name)
    request = (f"Audit the formulas in {name} with the xlsx skill's audit command and "
               f"confirm there are no formula errors.")
    return {"runner": "python", "code": code, "setup_runner": "python", "setup_code": setup,
            "artifact": name, "expect": {"kind": "tool", "json_status": False,
                                         "artifact": name}, "request": request, "mutation": None}


def gen_pivot(spec, rng):
    name, out = f"{spec.task_id}_wb.xlsx", f"{spec.task_id}_pivot.xlsx"
    setup = _book_setup(spec, rng, "regional revenue workbook for pivoting")
    code = _PIVOT.substitute(xlsx=_XLSX, name=name, out=out,
                             range="Data!A1:B5", metric="Revenue", dim="Region")
    request = (f"Build a PivotTable summary of {name} (rows: column A, values: column B) "
               f"with the xlsx skill's pivot command, writing {out}.")
    return {"runner": "python", "code": code, "setup_runner": "python", "setup_code": setup,
            "artifact": out, "expect": {"kind": "tool", "json_status": False,
                                        "artifact": out}, "request": request, "mutation": None}
