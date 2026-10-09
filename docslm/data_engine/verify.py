"""Real verification: execute scripts in the sandbox, then run the plugin's
own postcheck.py (docx) plus per-format structural checks (xlsx/pptx/pdf).

This module is the ground-truth provider for rejection sampling and for the
repair curriculum — no record is emitted without a measured verdict.
"""
import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from . import SANDBOX_ENV, PLUGIN_DIR

TIMEOUT = 90
PY = sys.executable  # run child scripts under the same interpreter as the build


def _env():
    env = dict(os.environ)
    env["NODE_PATH"] = str(SANDBOX_ENV / "node_modules")
    return env


def _sanitize(text: str, workdir: Path) -> str:
    """Strip machine-specific absolute paths — production diagnostics are relative."""
    return text.replace(str(workdir) + os.sep, "")


def _map_sandbox_paths(code: str) -> str:
    """Generated scripts reference plugin skills at the neutral mount /skills/;
    map to the real plugin dir at exec time (production sandboxes mount there)."""
    return code.replace("/skills/", str(PLUGIN_DIR) + "/")


def run_script(code: str, workdir: Path, runner: str, tag: str) -> dict:
    """Execute code in workdir; returns {'exit', 'stdout', 'stderr', 'seconds'}."""
    ext = "js" if runner == "node" else "py"
    path = workdir / f"{tag}.{ext}"
    path.write_text(_map_sandbox_paths(code), encoding="utf-8")
    cmd = ["node", path.name] if runner == "node" else [PY, path.name]
    import time
    t0 = time.time()
    try:
        proc = subprocess.run(cmd, cwd=workdir, env=_env(), timeout=TIMEOUT,
                              capture_output=True, text=True)
        result = {"exit": proc.returncode,
                  "stdout": _sanitize(proc.stdout[-2000:], workdir),
                  "stderr": _sanitize(proc.stderr[-2000:], workdir),
                  "seconds": round(time.time() - t0, 2)}
    except subprocess.TimeoutExpired:
        result = {"exit": -9, "stdout": "", "stderr": f"timeout after {TIMEOUT}s",
                  "seconds": TIMEOUT}
    result["script"] = tag + "." + ext
    return result


def postcheck_docx(docx_path: Path, workdir: Path | None = None) -> list[str]:
    """Run the plugin's own postcheck.py; returns violation strings."""
    cmd = [PY, str(PLUGIN_DIR / "docx" / "scripts" / "postcheck.py"), str(docx_path)]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
    except subprocess.TimeoutExpired:
        return ["postcheck:timeout"]
    out = (proc.stdout or "") + (proc.stderr or "")
    if workdir is not None:
        out = _sanitize(out, workdir)
    return _parse_postcheck(out, proc.returncode)


def _parse_postcheck(out: str, rc: int) -> list[str]:
    # postcheck.py prints a human-readable report; violations look like
    # "[FAIL] ..." / "✗ ..." lines or a JSON array. We accept both.
    violations = []
    try:
        start = out.find("[")
        end = out.rfind("]")
        if rc == 0 and start != -1 and end > start:
            candidate = json.loads(out[start:end + 1])
            if isinstance(candidate, list):
                for item in candidate:
                    if isinstance(item, dict):
                        status = str(item.get("status", "")).upper()
                        if status in ("FAIL", "ERROR", "WARN"):
                            violations.append(f"{item.get('check', 'check')}: {item.get('message', '')}")
                    elif isinstance(item, str):
                        violations.append(item)
                if candidate and not violations and not isinstance(candidate[0], (dict, str)):
                    pass
    except (json.JSONDecodeError, ValueError):
        pass
    if not violations:
        for line in out.splitlines():
            # item lines only: "[FAIL]/[ERROR]/[WARN]", ✗, "⚠ [check-name] …",
            # or any explicit violation text; never the "Passed x/y" summary line
            if re.search(r"\[(FAIL|ERROR|WARN)\]|✗|⚠.*\[|violation", line, re.I):
                violations.append(line.strip()[:220])
    if rc != 0 and not violations:
        violations.append(f"postcheck exited {rc}: {out.strip()[:220]}")
    return violations


# ---------- per-format structural checks ----------

_XLSX_REF_RE = re.compile(
    r"(?:([A-Za-z0-9_ \u4e00-\u9fff]+)!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")

def check_xlsx(path: Path, expect: dict) -> list[str]:
    from openpyxl import load_workbook
    v = []
    try:
        wb = load_workbook(path)
    except Exception as e:
        return [f"xlsx:load-failed:{e}"]
    sheet_names = {s.title for s in wb.worksheets}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "#REF!" in cell.value:
                    v.append(f"xlsx:REF-error@{ws.title}!{cell.coordinate}")
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    for m in _XLSX_REF_RE.finditer(cell.value):
                        ref_sheet = (m.group(1) or ws.title).strip().strip("'")
                        end_row = int(m.group(5) or m.group(3))
                        if ref_sheet not in sheet_names:
                            v.append(f"xlsx:unknown-sheet-ref@{ws.title}!{cell.coordinate}->{ref_sheet}")
                        elif ref_sheet in sheet_names and end_row > wb[ref_sheet].max_row + 64:
                            v.append(f"xlsx:out-of-range@{ws.title}!{cell.coordinate}->{ref_sheet}row{end_row}")
    data = wb["Data"] if "Data" in sheet_names else wb.worksheets[0]
    if expect.get("min_rows") and data.max_row < expect["min_rows"]:
        v.append(f"xlsx:too-few-rows:{data.max_row}<{expect['min_rows']}")
    if expect.get("charts"):
        n_charts = sum(len(ws._charts) for ws in wb.worksheets)
        if n_charts < expect["charts"]:
            v.append(f"xlsx:missing-charts:{n_charts}<{expect['charts']}")
    return v


def check_pptx(path: Path, expect: dict) -> list[str]:
    min_slides = expect.get("min_slides", 2)
    try:
        with zipfile.ZipFile(path) as z:
            slides = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
            bad = z.testzip()
    except Exception as e:
        return [f"pptx:zip-failed:{e}"]
    v = []
    if bad:
        v.append(f"pptx:corrupt-member:{bad}")
    if len(slides) < min_slides:
        v.append(f"pptx:too-few-slides:{len(slides)}<{min_slides}")
    try:
        from pptx import Presentation
        pres = Presentation(str(path))
        if len(pres.slides) != len(slides):
            v.append("pptx:slide-count-mismatch")
        for i, slide in enumerate(pres.slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip()]
            if not texts:
                v.append(f"pptx:empty-slide:{i}")
    except ImportError:
        pass  # python-pptx absent: zip-level check only
    except Exception as e:
        v.append(f"pptx:parse-failed:{e}")
    return v


def check_pdf(path: Path, expect: dict) -> list[str]:
    from pypdf import PdfReader
    v = []
    try:
        reader = PdfReader(str(path))
    except Exception as e:
        return [f"pdf:open-failed:{e}"]
    n = len(reader.pages)
    if n < expect.get("min_pages", 1):
        v.append(f"pdf:too-few-pages:{n}<{expect.get('min_pages')}")
    try:
        text = reader.pages[0].extract_text() or ""
        if len(text.strip()) < 5:
            v.append("pdf:empty-text-page1")
    except Exception as e:
        v.append(f"pdf:extract-failed:{e}")
    return v


def check_extract(path: Path, expect: dict) -> list[str]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception as e:
        return [f"extract:json-failed:{e}"]
    need = {"title", "paragraph_count", "body_sample", "tables"}
    missing = need - set(data)
    return [f"extract:missing-keys:{sorted(missing)}"] if missing else []


_TOOL_OK_STATUS = {"success", "ok", "passed", "pass"}


def check_tool(plan: dict, exec_result: dict, workdir: Path) -> list[str]:
    """Verify a plugin-CLI task: exit code, optional JSON status field,
    optional stdout sentinel, optional artifact existence."""
    expect = plan["expect"]
    v = []
    if exec_result["exit"] != 0:
        v.append(f"tool:exit{exec_result['exit']}: {exec_result['stderr'][-200:]}")
        return v
    stdout = exec_result.get("stdout", "")
    if expect.get("json_status"):
        status = None
        try:
            start = stdout.find("{")
            data = json.loads(stdout[start:stdout.rfind("}") + 1])
            status = data.get("status", data.get("pass"))
        except Exception:
            v.append("tool:stdout-not-json")
        if status is False or (isinstance(status, str) and status.lower() not in _TOOL_OK_STATUS):
            v.append(f"tool:status={status}")
    if expect.get("stdout_contains"):
        if expect["stdout_contains"] not in stdout:
            v.append(f"tool:sentinel-missing:{expect['stdout_contains']}")
    if expect.get("stdout_any_of") and not any(s in stdout for s in expect["stdout_any_of"]):
        v.append(f"tool:no-sentinel:{expect['stdout_any_of']}")
    if expect.get("artifact"):
        art = workdir / expect["artifact"]
        if not art.exists() or art.stat().st_size == 0:
            v.append(f"tool:artifact-missing:{expect['artifact']}")
    return v


CHECKERS = {"docx": None, "xlsx": check_xlsx, "pptx": check_pptx,
            "pdf": check_pdf, "extract": check_extract}


def verify(plan: dict, workdir: Path) -> dict:
    """Full verification of one plan: exec + format checks (+ plugin postcheck).

    Returns {'exec': result|None, 'artifact': bool, 'violations': [...], 'pass': bool}
    """
    out = {"exec": None, "artifact": False, "violations": [], "pass": False}
    if plan.get("setup_code"):
        sres = run_script(plan["setup_code"], workdir, plan["setup_runner"], "setup")
        if sres["exit"] != 0:
            out["violations"] = [f"setup failed: {sres['stderr'][-200:]}"]
            return out
    res = run_script(plan["code"], workdir, plan["runner"], "gen")
    out["exec"] = res
    artifact = workdir / plan["artifact"]
    if res["exit"] != 0:
        out["violations"] = [f"exec:exit{res['exit']}: {res['stderr'][-300:]}"]
        return out
    if not artifact.exists() or artifact.stat().st_size == 0:
        out["violations"] = ["artifact:missing-or-empty"]
        return out
    out["artifact"] = True
    kind = plan["expect"]["kind"]
    if kind == "tool":
        out["violations"] = check_tool(plan, res, workdir)
    elif kind == "docx":
        out["violations"] = postcheck_docx(artifact, workdir)
        for token in plan["expect"].get("ignore_warnings", []):
            out["violations"] = [v for v in out["violations"] if token not in v]
    else:
        out["violations"] = CHECKERS[kind](artifact, plan["expect"])
    out["violations"] = [_sanitize(v, workdir) if isinstance(v, str) else v
                         for v in out["violations"]]
    out["pass"] = not out["violations"]
    return out
