# DocSLM pptx inspect_edit — python-pptx edit
import json
from pptx import Presentation
from pptx.util import Pt

SRC, OUT = "task_000809_deck.pptx", "task_000809_upd.pptx"
prs = Presentation(SRC)
n_before = len(prs.slides)

slide = prs.slides.add_slide(prs.slide_layouts[len(prs.slide_layouts) - 1])
box = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(600), Pt(80))
tf = box.text_frame
tf.text = "Next Steps"
tf.paragraphs[0].runs[0].font.size = Pt(28)
tf.paragraphs[0].runs[0].font.bold = True

prs.save(OUT)
check = Presentation(OUT)
assert len(check.slides) == n_before + 1, "slide not appended"
texts = [sh.text_frame.text for s in check.slides for sh in s.shapes
         if sh.has_text_frame and sh.text_frame.text]
assert "Next Steps" in texts, "new slide title missing"
print(json.dumps({"exit": 0, "artifact": OUT, "slides": len(check.slides)}))
