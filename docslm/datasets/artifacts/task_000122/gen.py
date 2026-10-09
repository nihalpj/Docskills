# DocSLM format route — normalize spacing + heading styles
import json
from docx import Document

BASE, OUT = "task_000122_fmt.docx", "task_000122_fmt.docx"
doc = Document(BASE)

for p in doc.paragraphs:
    p.paragraph_format.line_spacing = 1.3
    if p.style is not None and p.style.name.startswith("Heading"):
        for r in p.runs:
            r.font.bold = True

def _fix(cell_or_doc):
    for p in cell_or_doc.paragraphs:
        p.paragraph_format.line_spacing = 1.3

for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            _fix(cell)
doc.save(OUT)
print(json.dumps({"exit": 0, "artifact": OUT, "formatted": True}))
