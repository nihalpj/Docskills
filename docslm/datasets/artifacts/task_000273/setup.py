# setup: attachment base document
from docx import Document
from docx.shared import Pt

doc = Document()
doc.add_heading("Risks and Challenges", level=1)
for t in ["Body paragraph one.", "Body paragraph two."]:
    p = doc.add_paragraph(t)
    if False:
        p.paragraph_format.line_spacing = Pt(12)   # too tight on purpose
doc.add_heading("Appendix Table", level=2)
for t in ["Table notes."]:
    p = doc.add_paragraph(t)
doc.add_table(rows=3, cols=3, style="Table Grid")
doc.save("task_000273_src.docx")
print("ok")
