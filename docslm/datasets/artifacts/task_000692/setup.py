# setup: attachment base document
from docx import Document
from docx.shared import Pt

doc = Document()
doc.add_heading("Department Quarterly Summary", level=1)
for t in "We recommend expanding channel investment and rebalancing resourcing to sustain growth. The business delivered steady growth this quarter, with core KPIs on target. Revenue concentration in the main product line increased and retention improved.":
    p = doc.add_paragraph(t)
    if True:
        p.paragraph_format.line_spacing = Pt(12)   # too tight on purpose
doc.add_heading("Data Analysis", level=2)
for t in ["Revenue concentration in the main product line increased and retention improved. We recommend expanding channel investment and rebalancing resourcing to sustain growth. Cost pressure and regional competition remain the principal risks to monitor."]:
    p = doc.add_paragraph(t)
doc.add_table(rows=3, cols=3, style="Table Grid")
doc.save("task_000692_fmt.docx")
print("ok")
