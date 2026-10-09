# DocSLM xlsx edit — extend workbook
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

wb = load_workbook("task_000622_src.xlsx")
ws = wb["Data"]
for r in [["Nimbus Storage", 356, 197], ["Pulse Dashboard", 131, 236], ["Nimbus Storage", 245, 163]]:
    ws.append(r)
last = ws.max_row
for row in ws.iter_rows(min_row=5 + 1, min_col=4, max_col=4):
    row[0].value = "=ROUND(B{0}-C{0},1)".format(row[0].row)
ws.cell(row=last + 1, column=1, value="Total")
ws.cell(row=last + 1, column=2, value="=SUM(B2:B{0})".format(last))
ws.cell(row=last + 1, column=3, value="=SUM(C2:C{0})".format(last))
ws.cell(row=last + 1, column=4, value="=ROUND(B{0}-C{0},1)".format(last + 1))
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill(fill_type="solid", fgColor="1F4E79")
wb.save("task_000622_upd.xlsx")
print(json.dumps({"exit": 0, "artifact": "task_000622_upd.xlsx", "appended": True}))
