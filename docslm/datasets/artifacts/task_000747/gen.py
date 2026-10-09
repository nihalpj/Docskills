# DocSLM xlsx pivot — PivotTable summary via the skill CLI
import json, subprocess, sys

proc = subprocess.run([sys.executable, "/home/nihal/.zcode/cli/plugins/cache/zcode-plugins-official/document-skills/0.1.4/skills/xlsx/xlsx.py", "pivot",
                       "--source", "Data!A1:B5", "--values", "Revenue", "--rows", "Region",
                       "task_000747_wb.xlsx", "task_000747_pivot.xlsx"],
                      capture_output=True, text=True, timeout=120)
assert proc.returncode == 0, f"pivot failed: {(proc.stdout + proc.stderr)[-300:]}"
from openpyxl import load_workbook
wb = load_workbook("task_000747_pivot.xlsx")
assert len(wb.sheetnames) >= 2, f"pivot output missing sheet: {wb.sheetnames}"
print(json.dumps({"exit": 0, "artifact": "task_000747_pivot.xlsx", "sheets": wb.sheetnames}))
