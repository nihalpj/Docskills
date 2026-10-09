# DocSLM read route — extract structure and content
import json
from docx import Document

doc = Document("task_000183_src.docx")
title = doc.paragraphs[0].text if doc.paragraphs else ""
body = [p.text for p in doc.paragraphs if p.text.strip()][1:11]
tables = [[[cell.text for cell in row.cells] for row in t.rows] for t in doc.tables]
out = {"title": title, "paragraph_count": len(doc.paragraphs),
       "body_sample": body, "tables": tables}
with open("task_000183_extract.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print(json.dumps({"exit": 0, "artifact": "task_000183_extract.json", "title": title}))
