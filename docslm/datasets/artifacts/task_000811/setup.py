# setup: attachment base document
from docx import Document
from docx.shared import Pt

doc = Document()
doc.add_heading("Project Basalt Interim Report", level=1)
for t in ["Progress on track; risks contained.", "Milestones M1 and M2 delivered on time."]:
    p = doc.add_paragraph(t)
    if True:
        p.paragraph_format.line_spacing = Pt(12)   # too tight on purpose
doc.add_heading("Milestone Review", level=2)
for t in ["Integration testing completed."]:
    p = doc.add_paragraph(t)
doc.add_table(rows=3, cols=3, style="Table Grid")
doc.save("task_000811_base.docx")
print("ok")
